"""create_document: host md/html rendering, sandbox pdf/docx/xlsx, naming, delivery; render_document.py."""
import asyncio
import importlib
import io
import json
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
from types import SimpleNamespace

import pytest

from agents import document_tools, web_tools
from agents.document_tools import (
    DocumentTools, has_h1, markdown_with_title, release_path, render_markdown_html, sanitize_filename, unique_path,
)

BANNER = "Secure container environment initialized\nExecuting command: python x\n----\n"
CONTENT = """# Quarterly report

Revenue grew **12%** with *strong* demand; see `tbl_1`.

## Sales by region

| Region | Q1 | Q2 |
|---|---:|---|
| North | 1200 | 1350.5 |
| South | 800 | n/a |

- keep investing
- hire two people

```python
print("code")
```
"""


class FakeSender:
    def __init__(self, fail=False):
        self.documents, self.fail = [], fail

    async def send_document(self, source, *, filename=None, caption=None):
        if self.fail:
            raise RuntimeError("chat not found")
        self.documents.append((str(source), filename, caption))


class FakeManager:
    """Plays the sandbox: records the call, 'renders' by writing bytes to the requested output."""

    def __init__(self, error=None, logs=None, raise_error=None):
        self.error, self.logs, self.raise_error, self.calls = error, logs, raise_error, []

    def run_script(self, user_id, script, args, timeout=90, network_enabled=False, mem_limit=None):
        opts, i = {}, 0
        while i < len(args):  # "--key value" pairs or "--key=value" tokens, as argparse reads them
            if "=" in args[i]:
                key, value = args[i].split("=", 1)
                opts[key], i = value, i + 1
            else:
                opts[args[i]], i = args[i + 1], i + 2
        source = Path("data") / user_id / opts["--input"]
        self.calls.append(SimpleNamespace(
            user_id=user_id, script=script, args=list(args), opts=opts, timeout=timeout,
            network_enabled=network_enabled, mem_limit=mem_limit,
            source_text=source.read_text(encoding="utf-8") if source.exists() else None,
        ))
        if self.raise_error:
            raise self.raise_error
        if self.logs is not None:
            return self.logs
        if self.error:
            return BANNER + json.dumps({"output": None, "bytes": 0, "error": self.error}) + "\n"
        out = Path("data") / user_id / opts["--output"]
        out.write_bytes(b"%FAKE " + opts["--format"].encode() + b" " * 2048)
        return BANNER + json.dumps({"output": opts["--output"], "bytes": out.stat().st_size, "error": None}) + "\n"


@pytest.fixture
def workspace(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    (tmp_path / "data" / "7").mkdir(parents=True)
    monkeypatch.setattr(web_tools, "_container_manager", None)
    return tmp_path


def use_manager(monkeypatch, manager):
    monkeypatch.setattr(web_tools, "_container_manager", manager)
    return manager


def make_tools(sender=None):
    return DocumentTools("7", sender if sender is not None else FakeSender())


# ---- helpers --------------------------------------------------------------------------------------
def test_sanitize_filename():
    assert sanitize_filename("Quarterly report: Q3/2026") == "Quarterly_report_Q3_2026"
    assert sanitize_filename("../../.hidden") == "hidden"
    assert sanitize_filename("  ") == "document" and sanitize_filename(None) == "document"
    assert sanitize_filename("reports/q3.pdf") == "reports_q3.pdf"
    assert len(sanitize_filename("x" * 500)) == document_tools.MAX_NAME_LENGTH
    assert sanitize_filename("Отчёт за квартал") == "Отчёт_за_квартал"  # unicode letters are file-name safe


def test_unique_path_never_overwrites(tmp_path):
    first = unique_path(tmp_path, "r", ".pdf")
    assert first == tmp_path / "r.pdf"
    (tmp_path / "r.pdf").write_bytes(b"1")
    release_path(first)
    second = unique_path(tmp_path, "r", ".pdf")
    assert second == tmp_path / "r-2.pdf"
    (tmp_path / "r-2.pdf").write_bytes(b"2")
    release_path(second)
    third = unique_path(tmp_path, "r", ".pdf")
    assert third == tmp_path / "r-3.pdf"
    release_path(third)


def test_unique_path_reserves_names_handed_out_but_not_written_yet(tmp_path):
    """Two create_document calls of one concurrent batch must not get the same name (check-then-create)."""
    a = unique_path(tmp_path, "r", ".pdf")
    b = unique_path(tmp_path, "r", ".pdf")
    assert (a, b) == (tmp_path / "r.pdf", tmp_path / "r-2.pdf") and not a.exists()
    release_path(a)
    assert unique_path(tmp_path, "r", ".pdf") == tmp_path / "r.pdf"  # released without being written: free again
    release_path(a)
    release_path(b)
    release_path(b)  # releasing twice is harmless


def test_has_h1_ignores_comments_inside_fenced_code():
    assert has_h1("# Title\n\ntext") and has_h1("intro\n\n# Title") and has_h1("# T")
    assert not has_h1("## only h2\n\ntext") and not has_h1("#hashtag") and not has_h1("")
    code = "Some text\n\n```python\n# a comment\nprint(1)\n```\n\nmore\n"
    assert not has_h1(code)
    assert markdown_with_title("Notes", code).startswith("# Notes\n\nSome text")
    assert "<h1>Notes</h1>" in render_markdown_html("Notes", code)
    assert has_h1(code + "\n# Real heading\n")


def test_render_markdown_html_tables_code_and_title():
    page = render_markdown_html("My title", CONTENT)
    assert page.startswith("<!DOCTYPE html>") and "<title>My title</title>" in page
    assert "<table>" in page and "<th" in page and "<td" in page and "1350.5" in page
    assert "<strong>12%</strong>" in page and "<em>strong</em>" in page and "<code>tbl_1</code>" in page
    assert '<pre><code class="language-python">' in page or "<pre><code" in page
    assert page.count("<h1") == 1  # the content's own H1 is used, the title is not repeated
    assert "<li>keep investing</li>" in page
    assert "border-collapse" in page  # the stylesheet is embedded

    untitled = render_markdown_html("Fallback <b>", "Just a paragraph.")
    assert "<h1>Fallback &lt;b&gt;</h1>" in untitled


# ---- host formats ---------------------------------------------------------------------------------
def test_md_is_written_sent_and_reported(workspace, monkeypatch):
    manager = use_manager(monkeypatch, FakeManager())
    sender = FakeSender()
    result = make_tools(sender).execute_tool("create_document", {"format": "md", "title": "Quarterly report", "content": CONTENT})

    path = workspace / "data" / "7" / "documents" / "Quarterly_report.md"
    assert path.read_text(encoding="utf-8") == CONTENT  # has its own H1: nothing prepended
    assert result == "Document saved to documents/Quarterly_report.md (1 KB) and sent to the user."
    # sent by its cwd-relative path (the bot runs with cwd=app/, like every other tool's sends)
    assert sender.documents == [("data/7/documents/Quarterly_report.md", "Quarterly_report.md", "Quarterly report")]
    assert manager.calls == []  # md never needs the sandbox


def test_md_without_a_heading_gets_the_title(workspace, monkeypatch):
    use_manager(monkeypatch, FakeManager())
    make_tools().execute_tool("create_document", {"format": "md", "title": "Notes", "content": "line one\n"})
    assert (workspace / "data" / "7" / "documents" / "Notes.md").read_text() == "# Notes\n\nline one\n"


def test_html_is_rendered_on_the_host(workspace, monkeypatch):
    use_manager(monkeypatch, FakeManager())
    result = make_tools().execute_tool("create_document", {"format": "html", "title": "Report", "content": CONTENT, "filename": "report.html"})
    page = (workspace / "data" / "7" / "documents" / "report.html").read_text(encoding="utf-8")
    assert "<table>" in page and "<style>" in page
    assert result.startswith("Document saved to documents/report.html (")


# ---- sandbox formats ----------------------------------------------------------------------------
def test_pdf_renders_html_in_the_sandbox_with_memory_and_no_network(workspace, monkeypatch):
    manager = use_manager(monkeypatch, FakeManager())
    sender = FakeSender()
    result = make_tools(sender).execute_tool("create_document", {"format": "pdf", "title": "Quarterly report", "content": CONTENT})

    call = manager.calls[0]
    assert call.script == "render_document.py" and call.network_enabled is False
    assert call.mem_limit == "1g" and call.timeout == 90
    assert call.opts["--format"] == "pdf" and call.opts["--title"] == "Quarterly report"
    assert "--title=Quarterly report" in call.args  # one token: a title starting with "-" is not an option
    assert call.opts["--output"] == "documents/Quarterly_report.pdf"
    assert call.opts["--input"].startswith("documents/.render_") and call.opts["--input"].endswith(".html")
    assert "<table>" in call.source_text and "<!DOCTYPE html>" in call.source_text  # pdf input is the rendered HTML
    assert not list((workspace / "data" / "7" / "documents").glob(".render_*"))  # temp source removed
    assert result == "Document saved to documents/Quarterly_report.pdf (2 KB) and sent to the user."
    assert sender.documents[0][1:] == ("Quarterly_report.pdf", "Quarterly report")


@pytest.mark.parametrize("fmt", ["docx", "xlsx"])
def test_docx_and_xlsx_get_the_markdown_source_and_default_memory(workspace, monkeypatch, fmt):
    manager = use_manager(monkeypatch, FakeManager())
    result = make_tools().execute_tool("create_document", {"format": fmt, "title": "T", "content": CONTENT})
    call = manager.calls[0]
    assert call.opts["--format"] == fmt and call.mem_limit is None and call.network_enabled is False
    assert call.opts["--input"].endswith(".md") and call.source_text == CONTENT
    assert result.startswith(f"Document saved to documents/T.{fmt} (")
    assert not list((workspace / "data" / "7" / "documents").glob(".render_*"))


def test_existing_names_are_kept_and_filenames_sanitized(workspace, monkeypatch):
    use_manager(monkeypatch, FakeManager())
    tools = make_tools()
    first = tools.execute_tool("create_document", {"format": "md", "title": "Plan", "content": "a"})
    second = tools.execute_tool("create_document", {"format": "md", "title": "Plan", "content": "b"})
    assert "documents/Plan.md" in first and "documents/Plan-2.md" in second
    assert (workspace / "data" / "7" / "documents" / "Plan.md").read_text() == "# Plan\n\na"

    third = tools.execute_tool("create_document", {"format": "pdf", "title": "x", "content": "a", "filename": "../My Plan: v2.PDF"})
    assert "documents/My_Plan_v2.pdf" in third


def test_sandbox_failures_become_error_strings(workspace, monkeypatch):
    use_manager(monkeypatch, FakeManager(error="TimeoutError: Timed out after 75 s while rendering the pdf"))
    assert make_tools().execute_tool("create_document", {"format": "pdf", "title": "T", "content": "a"}) == \
        "Error: TimeoutError: Timed out after 75 s while rendering the pdf"

    use_manager(monkeypatch, FakeManager(logs=BANNER + "Error (exit code 137):\nKilled\n"))
    result = make_tools().execute_tool("create_document", {"format": "docx", "title": "T", "content": "a"})
    assert result.startswith("Error: the sandbox renderer returned no result") and "Killed" in result

    use_manager(monkeypatch, FakeManager(raise_error=RuntimeError("no docker")))
    result = make_tools().execute_tool("create_document", {"format": "xlsx", "title": "T", "content": "a"})
    assert result.startswith("Error: could not start the sandbox renderer") and "no docker" in result
    assert not list((workspace / "data" / "7" / "documents").glob(".render_*"))  # cleaned up on every path


def test_a_title_starting_with_a_dash_reaches_the_renderer(workspace, monkeypatch, rd):
    manager = use_manager(monkeypatch, FakeManager())
    make_tools().execute_tool("create_document", {"format": "docx", "title": "-Draft v2", "content": "a"})
    assert manager.calls[0].opts["--title"] == "-Draft v2"
    parsed = rd.parse_args(manager.calls[0].args)  # the sandbox script's own argparse accepts it as sent
    assert parsed.title == "-Draft v2" and parsed.format == "docx"


def test_concurrent_calls_with_the_same_title_get_different_files(workspace, monkeypatch):
    use_manager(monkeypatch, FakeManager())
    tools = DocumentTools("7", None)
    with ThreadPoolExecutor(max_workers=4) as pool:
        results = list(pool.map(lambda n: tools.create_document("md", "Same", f"body {n}"), range(4)))
    names = sorted(r.split("documents/")[1].split(" ")[0] for r in results)
    assert names == ["Same-2.md", "Same-3.md", "Same-4.md", "Same.md"]
    assert sorted(p.name for p in (workspace / "data" / "7" / "documents").glob("Same*.md")) == names
    assert document_tools._reserved == set()  # every name released after its write


def test_bad_arguments(workspace, monkeypatch):
    manager = use_manager(monkeypatch, FakeManager())
    tools = make_tools()
    assert tools.execute_tool("create_document", {"format": "pptx", "title": "T", "content": "a"}).startswith("Error: unsupported format 'pptx'")
    assert tools.execute_tool("create_document", {"format": "md", "title": "T", "content": "  "}) == "Error: content is empty"
    assert tools.execute_tool("nope", {}) == "Unknown tool: nope"
    assert manager.calls == []
    assert [t["name"] for t in tools.tools_schema] == ["create_document"]
    assert tools.tools_schema[0]["input_schema"]["properties"]["format"]["enum"] == ["md", "html", "pdf", "docx", "xlsx"]


def test_delivery_without_a_sender_or_with_a_failing_one(workspace, monkeypatch):
    use_manager(monkeypatch, FakeManager())
    assert DocumentTools("7", None).execute_tool("create_document", {"format": "md", "title": "T", "content": "a"}) == \
        "Document saved to documents/T.md (1 KB) (no chat to send it to)."
    result = DocumentTools("7", FakeSender(fail=True)).execute_tool("create_document", {"format": "md", "title": "T", "content": "a"})
    assert result.startswith("Document saved to documents/T-2.md (1 KB), but sending it failed: chat not found")
    assert "send_file_path_to_user_via_telegram" in result


async def test_sending_from_the_worker_thread_uses_the_bots_loop(workspace, monkeypatch):
    """Registered blocking=True: execute_tool runs in asyncio.to_thread and must still reach the sender."""
    use_manager(monkeypatch, FakeManager())
    sender = FakeSender()
    tools = DocumentTools("7", sender)  # created on the loop, like ChainOfThoughtAgent does
    assert tools._loop is asyncio.get_running_loop()
    result = await asyncio.to_thread(tools.execute_tool, "create_document", {"format": "pdf", "title": "T", "content": "a"})
    assert result.endswith("and sent to the user.") and len(sender.documents) == 1
    # calling the sync tool on the loop thread itself would deadlock: it is refused instead
    assert "must run in a worker thread" in tools.execute_tool("create_document", {"format": "md", "title": "T", "content": "a"})


# ---- render_document.py (the sandbox side, importable here) ------------------------------------
@pytest.fixture
def rd():
    return importlib.import_module("secure_container.scripts.render_document")


def test_parse_markdown_blocks(rd):
    blocks = rd.parse_markdown(CONTENT)
    assert [b["type"] for b in blocks] == ["heading", "paragraph", "heading", "table", "list", "code"]
    assert blocks[0] == {"type": "heading", "level": 1, "text": "Quarterly report"}
    assert blocks[3]["header"] == ["Region", "Q1", "Q2"] and blocks[3]["rows"] == [["North", "1200", "1350.5"], ["South", "800", "n/a"]]
    assert blocks[4] == {"type": "list", "ordered": False, "items": ["keep investing", "hire two people"]}
    assert blocks[5] == {"type": "code", "lang": "python", "text": 'print("code")'}
    numbered = rd.parse_markdown("1. a\n2. b\n\nText\nmore text\n")
    assert numbered == [{"type": "list", "ordered": True, "items": ["a", "b"]}, {"type": "paragraph", "text": "Text more text"}]
    # ragged table rows are padded or cut to the header width; a lone pipe line is a paragraph
    ragged = rd.parse_markdown("a | b\n--|--\n1 |\n2 | 3 | 4\nno pipe ends the table\n")
    assert ragged[0]["rows"] == [["1", ""], ["2", "3"]] and ragged[1]["type"] == "paragraph"
    assert rd.parse_markdown("just | a pipe")[0]["type"] == "paragraph"
    assert [b["type"] for b in rd.parse_markdown("a | b\n---\ntext")] == ["paragraph"]  # a rule is not a 2-column delimiter


def test_inline_runs_and_cell_values(rd):
    assert rd.inline_runs("a **b** *c* `d` _e_") == [("a ", ""), ("b", "bold"), (" ", ""), ("c", "italic"), (" ", ""), ("d", "code"), (" _e_", "")]
    assert rd.plain_text("**bold** and `code`") == "bold and code"
    assert rd.inline_runs("*x*, (*y*) and **z**.") == [("x", "italic"), (", (", ""), ("y", "italic"), (") and ", ""), ("z", "bold"), (".", "")]
    assert rd.cell_value("1200") == 1200 and rd.cell_value(" 1350.5 ") == 1350.5 and rd.cell_value("-3") == -3
    assert rd.cell_value("n/a") == "n/a" and rd.cell_value("**12**") == 12 and rd.cell_value("1e3") == 1000.0
    assert rd.cell_value("") == ""


def test_identifiers_that_look_numeric_stay_text_in_xlsx(rd):
    """Zip codes, phone numbers and long ids lost their leading zeros, their "+" or their last digits."""
    for text in ("02134", "0791234567", "007", "-007", "+44", "+447911123456", "12345678901234567890",
                 "4111111111111111", "00.5"):
        assert rd.cell_value(text) == text, text
    # ordinary quantities still become numbers
    assert rd.cell_value("0") == 0 and rd.cell_value("0.5") == 0.5 and rd.cell_value("-0.25") == -0.25
    assert rd.cell_value("123456789012345") == 123456789012345  # 15 digits: exact in Excel
    assert rd.cell_value("1e3") == 1000.0 and rd.cell_value("0.000123") == 0.000123
    wb = rd.build_xlsx("| Phone | ZIP | Order | Qty |\n|---|---|---|---|\n| 0791234567 | 02134 | 12345678901234567890 | 3 |\n", "T")
    assert [[c.value for c in r] for r in wb["Table 1"].iter_rows()][1] == ["0791234567", "02134", "12345678901234567890", 3]


def test_markers_that_are_not_markup_survive_untouched(rd):
    """snake_case identifiers and arithmetic are everyday report content, not italics."""
    for text in ("user_id_hash", "column_name_x", "2*3*4", "a*b", "_private", "__init__", "foo_ bar_", "5 * 3 = 15", "**"):
        assert rd.plain_text(text) == text, text
        assert rd.cell_value(text) == text, text
    assert rd.inline_runs("user_id_hash") == [("user_id_hash", "")]
    assert rd.plain_text("`user_id` is *italic*") == "user_id is italic"
    wb = rd.build_xlsx("| column_name_x | total_count |\n|---|---|\n| user_id_hash | 2*3*4 |\n", "T")
    assert [[c.value for c in r] for r in wb["Table 1"].iter_rows()] == [["column_name_x", "total_count"], ["user_id_hash", "2*3*4"]]
    doc = rd.build_docx("# head_ing\n\nuser_id_hash and 2*3*4\n", "")
    assert [p.text for p in doc.paragraphs] == ["head_ing", "user_id_hash and 2*3*4"]


def test_build_docx(rd):
    from docx import Document

    buffer = io.BytesIO()
    rd.build_docx(CONTENT, "Quarterly report").save(buffer)
    doc = Document(io.BytesIO(buffer.getvalue()))
    styles = [p.style.name for p in doc.paragraphs]
    texts = [p.text for p in doc.paragraphs]
    assert "Heading 1" in styles and "Heading 2" in styles and styles.count("Title") == 0  # content has its own H1
    assert texts[0] == "Quarterly report"
    body = next(p for p in doc.paragraphs if p.text.startswith("Revenue grew"))
    assert [(r.text, bool(r.bold), bool(r.italic)) for r in body.runs][:4] == [
        ("Revenue grew ", False, False), ("12%", True, False), (" with ", False, False), ("strong", False, True)]
    assert any(r.font.name == "Courier New" for r in body.runs)
    assert styles.count("List Bullet") == 2
    code = next(p for p in doc.paragraphs if p.text == 'print("code")')
    assert code.runs[0].font.name == "Courier New"
    assert len(doc.tables) == 1
    table = doc.tables[0]
    assert [c.text for c in table.rows[0].cells] == ["Region", "Q1", "Q2"]
    assert all(r.bold for c in table.rows[0].cells for r in c.paragraphs[0].runs)
    assert [c.text for c in table.rows[2].cells] == ["South", "800", "n/a"]
    assert doc.core_properties.title == "Quarterly report"

    untitled = rd.build_docx("plain text only", "Given title")
    assert [p.style.name for p in untitled.paragraphs][0] == "Title" and untitled.paragraphs[0].text == "Given title"


def test_build_xlsx(rd):
    wb = rd.build_xlsx(CONTENT + "\n## Staff\n\n```csv\nname,age\nAnn,31\nBob,40\n```\n\n## Staff\n\n| a |\n|---|\n| 1 |\n", "T")
    assert wb.sheetnames == ["Sales by region", "Staff", "Staff (2)"]
    rows = [[c.value for c in r] for r in wb["Sales by region"].iter_rows()]
    assert rows == [["Region", "Q1", "Q2"], ["North", 1200, 1350.5], ["South", 800, "n/a"]]
    assert all(c.font.bold for c in wb["Sales by region"][1])
    assert [[c.value for c in r] for r in wb["Staff"].iter_rows()] == [["name", "age"], ["Ann", 31], ["Bob", 40]]
    assert wb.properties.title == "T"

    # a cell starting with "=" is table text, not a formula (openpyxl would store it as one)
    from openpyxl import load_workbook

    formulas = rd.build_xlsx("| step | note |\n|---|---|\n| =SUM(1) | => next step |\n| =1+1 | ok |\n", "T")
    cells = [c for row in formulas["Table 1"].iter_rows(min_row=2) for c in row]
    assert [(c.value, c.data_type) for c in cells] == [("=SUM(1)", "s"), ("=> next step", "s"), ("=1+1", "s"), ("ok", "s")]
    buffer = io.BytesIO()
    formulas.save(buffer)
    reloaded = load_workbook(io.BytesIO(buffer.getvalue()))["Table 1"]
    assert [(c.value, c.data_type) for c in reloaded[2]] == [("=SUM(1)", "s"), ("=> next step", "s")]

    plain = rd.build_xlsx("no tables here\nsecond line", "Notes: [draft]/v1")
    assert plain.sheetnames == ["Notes draft v1"]  # Excel forbids []:*?/\ in sheet names
    assert [[c.value for c in r] for r in plain["Notes draft v1"].iter_rows()] == [["no tables here"], ["second line"]]
    assert rd.sheet_name("x" * 40, set()) == "x" * 31


def test_render_document_main_writes_the_file_and_prints_json(rd, tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(rd, "_printed", False)
    source = tmp_path / "in.md"
    source.write_text(CONTENT, encoding="utf-8")
    out = tmp_path / "sub" / "out.xlsx"
    assert rd.main(["--format", "xlsx", "--input", str(source), "--output", str(out), "--title=-T\u2028x"]) == 0
    printed = capsys.readouterr().out
    assert printed.isascii()  # line separators in a title or path never split the JSON line
    data = json.loads(printed.strip().splitlines()[-1])
    assert data == {"output": str(out), "bytes": out.stat().st_size, "error": None} and out.stat().st_size > 0

    monkeypatch.setattr(rd, "_printed", False)
    assert rd.main(["--format", "docx", "--input", str(tmp_path / "missing.md"), "--output", str(tmp_path / "o.docx")]) == 0
    data = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert data["output"] is None and data["bytes"] == 0 and "FileNotFoundError" in data["error"]


def test_render_document_pdf_path_needs_the_browser_and_reports_its_absence(rd, tmp_path, capsys, monkeypatch):
    monkeypatch.setattr(rd, "_printed", False)
    monkeypatch.setitem(__import__("sys").modules, "playwright", None)
    monkeypatch.setitem(__import__("sys").modules, "playwright.sync_api", None)
    source = tmp_path / "in.html"
    source.write_text("<html><body>x</body></html>")
    assert rd.main(["--format", "pdf", "--input", str(source), "--output", str(tmp_path / "o.pdf")]) == 0
    data = json.loads(capsys.readouterr().out.strip().splitlines()[-1])
    assert data["error"] and data["output"] is None
